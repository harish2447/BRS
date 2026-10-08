from pathlib import Path
import re
import sys
import uuid
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

BASE_DIR = Path(__file__).resolve().parent
PROMPT_PATH = BASE_DIR / "BRS Generation Prompt.txt"
DEFAULT_TOM = BASE_DIR / "ImagineR3.7_SSJSFJ_J5and6_Oct Release_TOM_V0.9.xlsx"

HEADER_ROW = [
    "Ref #",
    "Requirement Category",
    "Requirement Title",
    "Requirement Description",
    "Acceptance Criteria",
    "Requirement SME",
    "MoSCoW",
    "System",
    "Reference to Other Related Documents If Applicable",
]


def load_prompt(path: Path) -> str:
    if not path.exists():
        raise FileNotFoundError(f"Prompt file not found: {path}")
    return path.read_text(encoding="utf-8", errors="replace")


def infer_journey_name(tom_path: Path) -> str:
    name = tom_path.stem.replace("_TOM", "")
    name = re.sub(r"_V\d+(?:\.\d+)?$", "", name)
    name = name.replace("_", " ")
    return name.strip()


def find_reference_template(base_dir: Path):
    candidates = [
        base_dir / "Updated BRS 5&6 23_09.xlsx",
        base_dir / "Updated BRS 5&6.xlsx",
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    return None


def normalize_text(value):
    if value is None:
        return ""
    text = str(value).strip()
    return re.sub(r"\s+", " ", text)


def text_from_rows(ws):
    pieces = []
    for row in ws.iter_rows(values_only=True):
        values = [normalize_text(v) for v in row if normalize_text(v)]
        if values:
            pieces.append(" ".join(values))
    return " ".join(pieces)


def get_tom_text(wb):
    all_text = []
    for ws in wb.worksheets:
        text = text_from_rows(ws)
        if text:
            all_text.append(text)
    return "\n".join(all_text)


def append_unique(req_list, item):
    key = re.sub(r"\s+", " ", item["title"]).lower()
    for existing in req_list:
        if re.sub(r"\s+", " ", existing["title"]).lower() == key:
            return
    req_list.append(item)


def build_requirement_rows(wb):
    tom_text = get_tom_text(wb)
    low = tom_text.lower()
    requirements = []

    requirement_templates = [
        {
            "ref": "FR001",
            "category": "Quote - Coverage and address search",
            "title": "Retrieve NBN Business plan only check box functionality in Quote address Search",
            "keywords": ["business plan only", "retrieve only business plan option", "nbn business-plan-only"],
            "description": "The solution must allow authorised users to request NBN business plans only when creating an eligible business quote, while preventing consumer users from using the option.",
            "criteria": "AC1: For a new business customer, the user can select the NBN business-plan-only option before performing the address search.\nAC2: When the option is selected, the system sends a request for NBN coverage and displays only eligible NBN business plans on the product-list page.\nAC3: For a consumer customer, the NBN business-plan-only option is unavailable and cannot be selected.\nAC4: For an existing business customer, the option is available according to the customer eligibility rules.",
            "system": "Siebel",
        },
        {
            "ref": "FR002",
            "category": "Quote - Coverage and address search",
            "title": "Address Search and One SQ request in Quote address Search",
            "keywords": ["psma", "manual address", "location id", "onesq request"],
            "description": "The solution must accept the supported address-search methods and send the appropriate OneSQ request for the selected address input.",
            "criteria": "AC1: The user can search by PSMA address, manual address, NBN location ID, or an address with multiple location IDs.\nAC2: A PSMA or manual-address search sends one request containing the configured NBN, Vision, Opticomm, and Mapshed networks.\nAC3: A location-ID search sends a network-specific NBN request and returns the NBN response for the selected location.\nAC4: A multiple-location search sends the required requests for the available locations and associates each response with the correct location ID.\nAC5: The address-search results display the response returned for the selected search method.",
            "system": "Siebel",
        },
        {
            "ref": "FR003",
            "category": "Quote - Coverage and address search",
            "title": "Mobile Coverage as per OneSQ response in Quote Coverage Check",
            "keywords": ["4g", "5g", "mobile coverage"],
            "description": "The solution must display the mobile coverage results returned by OneSQ during the quote coverage check.",
            "criteria": "AC1: The coverage page displays 4G as Available or Not available according to the OneSQ response.\nAC2: The coverage page displays 5G outdoor and indoor availability according to the OneSQ response.\nAC3: The coverage page displays 5G NSA as Available or Not available according to the OneSQ response.",
            "system": "Siebel",
        },
        {
            "ref": "FR004",
            "category": "Quote - Coverage and address search",
            "title": "Fixed Wholesaler Coverage as per OneSQ response in Quote Coverage Check",
            "keywords": ["wholesaler", "preferred wholesaler", "availability as per onesq response"],
            "description": "The solution must display fixed-wholesaler availability and identify the preferred wholesaler using the OneSQ response and configured priority.",
            "criteria": "AC1: The coverage page displays the availability status for NBN, Vision, and Opticomm according to the OneSQ response.\nAC2: The system selects the preferred available wholesaler using the configured wholesaler priority.\nAC3: When more than one wholesaler is available, an authorised user can select an alternative wholesaler from the available list.\nAC4: The wholesaler override action is hidden from users without the required responsibility.\nAC5: The selected wholesaler is carried forward to the next quote step.",
            "system": "Siebel",
        },
        {
            "ref": "FR005",
            "category": "Quote - Coverage and address search",
            "title": "Fixed Wholesaler Technology as per OneSQ response in Quote Coverage Check",
            "keywords": ["technology type", "service class", "maximum attainable speed", "development charge"],
            "description": "The solution must display the technology, service, speed, upgrade, and capacity information returned for each available fixed wholesaler.",
            "criteria": "AC1: The system displays the location ID, service class, technology type, maximum attainable speed, development status, and fibre-upgrade status returned for NBN.\nAC2: The system displays the corresponding location, technology, speed, and upgrade information returned for Vision and Opticomm.\nAC3: For FTTP or HFC technology, the system displays active services, total capacity, used capacity, and remaining capacity.\nAC4: For copper technology, NTD capacity information is not displayed.\nAC5: When the response identifies a new development, the applicable development-charge message is displayed.",
            "system": "Siebel",
        },
        {
            "ref": "FR006",
            "category": "Quote - Coverage and address search",
            "title": "Fixed Wireless Coverage as per OneSQ response in Quote Coverage Check",
            "keywords": ["fixed wireless", "4g and 5g", "coverage override"],
            "description": "The solution must display 4G and 5G fixed-wireless availability and provide the authorised coverage override action when coverage is unavailable.",
            "criteria": "AC1: The coverage page displays 4G and 5G NSA fixed-wireless availability according to the OneSQ response.\nAC2: When required coverage is unavailable, an authorised user can access the coverage override action and provide the override user ID.\nAC3: After an authorised override, the eligible 4G and 5G fixed-wireless propositions are displayed on the next product-list page.\nAC4: The coverage override action is unavailable to users without the required responsibility.",
            "system": "Siebel",
        },
        {
            "ref": "FR007",
            "category": "Quote - Coverage and address search",
            "title": "Add PLP Page Navigation after coverage check in Quote",
            "keywords": ["next on the coverage page", "product-list page", "mobile tab"],
            "description": "The solution must take the user to the appropriate product-list page after the coverage check is completed.",
            "criteria": "AC1: Selecting Next on the coverage page opens the product-list page.\nAC2: The Mobile tab is selected by default when the product-list page opens.\nAC3: The user can select the Fixed tab to continue the fixed-service quote journey.",
            "system": "Siebel",
        },
        {
            "ref": "FR008",
            "category": "Quote - Coverage and address search",
            "title": "Coverage Details and One SQ request in Quote Add PLP Page",
            "keywords": ["coverage check result", "preferred wholesaler", "override wholesaler"],
            "description": "The solution must carry the address coverage response into the fixed product-list page and allow the user to review or change the address.",
            "criteria": "AC1: The system shall ensure that coverage check result is displayed under Coverage check section in Fixed PLP page.\nAC2: Option to change address is available.\nAC3: OneSQ request should be triggered in the backend and coverage results is displayed the same address-based coverage results, preferred wholesaler, and available coverage actions shown in the current coverage response along with preferred wholesaler, override wholesaler button and overrider coverage check button.",
            "system": "Siebel",
        },
        {
            "ref": "FR009",
            "category": "Quote - Product and plan selection",
            "title": "Product sections available in Quote Add PLP Page",
            "keywords": ["fixed plans and devices", "accessories", "other devices"],
            "description": "The solution must provide separate product sections for fixed plans and devices, accessories, and other devices.",
            "criteria": "AC1: The product-list page displays the Fixed plans and devices section.\nAC2: The product-list page displays the Accessories section.\nAC3: The product-list page displays the Other devices section.\nAC4: Each section displays only the products applicable to the selected quote and coverage result.",
            "system": "Siebel",
        },
        {
            "ref": "FR010",
            "category": "Quote - Product and plan selection",
            "title": "Proposition filteration logic in Quote Add PLP Page",
            "keywords": ["proposition", "filteration", "preferred wholesaler"],
            "description": "The solution must show only propositions that match the selected wholesaler, customer plan type, and available fixed-wireless coverage.",
            "criteria": "AC1: The proposition list contains only propositions applicable to the selected preferred wholesaler and fixed-wireless coverage.\nAC2: When the NBN business-plan-only option is selected, only eligible business propositions and plans are displayed.\nAC3: The user can search for propositions associated with the selected wholesaler.\nAC4: After a coverage override, the eligible 4G and 5G fixed-wireless propositions are displayed even when coverage is unavailable.",
            "system": "Siebel",
        },
        {
            "ref": "FR011",
            "category": "Quote - Product and plan selection",
            "title": "Plan filteration logic in Quote Add PLP Page",
            "keywords": ["plan list", "estimated maximum attainable speed", "tier above"],
            "description": "The solution must filter available plans using the wholesaler, proposition, technology type, and estimated maximum attainable speed.",
            "criteria": "AC1: The plan list is filtered using the selected wholesaler, proposition, technology type, and estimated maximum attainable speed.\nAC2: For copper technologies, the plan list follows the estimated speed range and displays the applicable warning on the plan tile.\nAC3: The user can select a plan up to one speed tier above the highest estimated download-speed value where the rule permits it.\nAC4: For fibre and HFC technologies, the plan list is filtered using the applicable maximum attainable speed.",
            "system": "Siebel",
        },
        {
            "ref": "FR012",
            "category": "Quote - Product and plan selection",
            "title": "Estimated MAS Warning logic in Quote Add PLP Page",
            "keywords": ["estimated mas warning", "speed-limitation message"],
            "description": "The solution must warn the user when a selected plan may exceed the estimated maximum attainable speed for the address.",
            "criteria": "AC1: The plan tile displays the configured speed-limitation message and amber warning icon when the estimated speed rule applies.\nAC2: The system applies the estimated-speed rule separately for the supported copper, fibre, HFC, Vision, and Opticomm technologies.\nAC3: Plans below the applicable estimated-speed threshold are displayed without an incorrect warning or are excluded according to the configured rule.",
            "system": "Siebel",
        },
        {
            "ref": "FR013",
            "category": "Quote - Product and plan selection",
            "title": "Show Current Plan toggle in Plan section",
            "keywords": ["show current plan"],
            "description": "The solution must control the visibility of the current-plan option according to the customer connection type.",
            "criteria": "AC1: The Show Current Plan toggle is unavailable for a new-connect journey.\nAC2: When the toggle is available for an eligible journey, selecting it displays the customer current plan.",
            "system": "Siebel",
        },
        {
            "ref": "FR014",
            "category": "Quote - Product and plan selection",
            "title": "Check Delivery estimate in modem Section",
            "keywords": ["check delivery estimate", "vhagetestimateddates", "modem"],
            "description": "The solution must allow the user to request an estimated modem delivery date using a suburb or postcode.",
            "criteria": "AC1: The system shall ensure that Check delivery estimate is optional.\nAC2: Agent can search suburb or post code to find the matching list of suburbs shown as inline list.\nAC3: Once a suburb is selected, system will show the combination of suburb state and postcode.\nAC4: The system shall ensure that Once a suburb is selected, system should trigger the API: VHAGetEstimatedDates by passing the three/active devices shown in the UI to get the stock status and the estimated delivery date.\nAC5: The system shall ensure that coverage check checkbox should be enabled when coverage check result / address is available in the session, Agent can select this checkbox, which will pre-populate the suburb state and postcode in the check delivery estimate field and trigger the API: VHAGetEstimatedDates by passing the three/active devices shown in the UI to get the stock status and the estimated delivery date.\nAC6: Based on the ERP API responses, in-stock, out-of-stock, back- Order, low stock lozenge is displayed.",
            "system": "Siebel",
        },
        {
            "ref": "FR015",
            "category": "Quote - Product and plan selection",
            "title": "Modem Filteration in modem Section",
            "keywords": ["modem list", "compatible modems", "sort by"],
            "description": "The solution must show only compatible modems and allow the user to sort the available modem results.",
            "criteria": "AC1: The modem list contains only modems compatible with the selected wholesaler and technology.\nAC2: The user can sort the modem results using the available Sort by options.",
            "system": "Siebel",
        },
        {
            "ref": "FR016",
            "category": "Quote - Product and plan selection",
            "title": "Show current Modem toggle functionality in Modem Section",
            "keywords": ["show current modem"],
            "description": "The solution must control the visibility of the current-modem option according to the customer connection type.",
            "criteria": "AC1: The Show Current Modem toggle is unavailable for a new-connect journey.\nAC2: When the toggle is available for an eligible journey, selecting it displays the customer current modem.",
            "system": "Siebel",
        },
        {
            "ref": "FR017",
            "category": "Quote - Product and plan selection",
            "title": "Check Delivery estimate in Accessories Section",
            "keywords": ["accessories", "check delivery estimate"],
            "description": "The solution must allow the user to request an estimated accessory delivery date using a suburb or postcode.",
            "criteria": "AC1: The system shall ensure that Check delivery estimate is optional.\nAC2: Agent can search suburb or post code to find the matching list of suburbs shown as inline list.\nAC3: Once a suburb is selected, system will show the combination of suburb state and postcode.\nAC4: The system shall ensure that Once a suburb is selected, system should trigger the API: VHAGetEstimatedDates by passing the three/active devices shown in the UI to get the stock status and the estimated delivery date.\nAC5: The system shall ensure that coverage check checkbox should be enabled when coverage check result / address is available in the session, Agent can select this checkbox, which will pre-populate the suburb state and postcode in the check delivery estimate field and trigger the API: VHAGetEstimatedDates by passing the three/active devices shown in the UI to get the stock status and the estimated delivery date.\nAC6: Based on the ERP API responses, in-stock, out-of-stock, back- Order, low stock lozenge is displayed.",
            "system": "Siebel",
        },
        {
            "ref": "FR018",
            "category": "Quote - Product and plan selection",
            "title": "Accessories Filteration in Accessories Section",
            "keywords": ["accessories list", "vendor or brand", "sort by"],
            "description": "The solution must filter accessories by the selected vendor or brand and allow the results to be sorted.",
            "criteria": "AC1: The accessories list is filtered using the selected vendor or brand.\nAC2: The user can sort the accessory results using the available Sort by options.",
            "system": "Siebel",
        },
        {
            "ref": "FR019",
            "category": "Quote - Product and plan selection",
            "title": "Show Current APPs toggle in Plan section",
            "keywords": ["show current apps"],
            "description": "The solution must control the visibility of the current-apps option according to the customer connection type.",
            "criteria": "AC1: The Show Current APPs toggle is unavailable for a new-connect journey.\nAC2: When the toggle is available for an eligible journey, selecting it displays the customer current applications.",
            "system": "Siebel",
        },
        {
            "ref": "FR020",
            "category": "Quote - Product and plan selection",
            "title": "Check Delivery estimate in Other devices Section",
            "keywords": ["other devices", "check delivery estimate"],
            "description": "The solution must allow the user to request an estimated delivery date for other devices using a suburb or postcode.",
            "criteria": "AC1: The system shall ensure that Check delivery estimate is optional.\nAC2: Agent can search suburb or post code to find the matching list of suburbs shown as inline list.\nAC3: Once a suburb is selected, system will show the combination of suburb state and postcode.\nAC4: The system shall ensure that Once a suburb is selected, system should trigger the API: VHAGetEstimatedDates by passing the three/active devices shown in the UI to get the stock status and the estimated delivery date.\nAC5: The system shall ensure that coverage check checkbox should be enabled when coverage check result / address is available in the session, Agent can select this checkbox, which will pre-populate the suburb state and postcode in the check delivery estimate field and trigger the API: VHAGetEstimatedDates by passing the three/active devices shown in the UI to get the stock status and the estimated delivery date.\nAC6: Based on the ERP API responses, in-stock, out-of-stock, back- Order, low stock lozenge is displayed.",
            "system": "Siebel",
        },
        {
            "ref": "FR021",
            "category": "Quote - Product and plan selection",
            "title": "Other devices Filteration in Other devices Section",
            "keywords": ["show fixed devices", "sort by", "arcadyan mesh"],
            "description": "The solution must enable the user to filter other devices and display only the relevant device options for the selected journey.",
            "criteria": "AC1: The system shall ensure that 'Show fixed devices' checkbox is enabled by default to display Arcadyan Mesh secondary devices.\nAC2: The system shall ensure that unchecking the ' Show fixed devices' checkbox should enable user to select other devices as well other than Mesh.\nAC3: The system shall ensure that ' Sort by' option is working as per selection.\nAC4: The system shall ensure that 'Show fixed devices' checkbox is enabled by default.\nAC5: But Arcadyan mesh should not be displayed for Vision.\nAC6: But Arcadyan mesh should not be displayed for FWA.",
            "system": "Siebel",
        },
        {
            "ref": "FR022",
            "category": "Quote - Product and plan selection",
            "title": "Show Current Devices toggle in Plan section",
            "keywords": ["show current devices"],
            "description": "The solution must enable the user to show Current Devices toggle in Plan section.",
            "criteria": "AC1: The system shall ensure that Toggle for 'Show Current Devices' should not be available for new connect.",
            "system": "Siebel",
        },
        {
            "ref": "FR023",
            "category": "Quote - Cart and quote creation",
            "title": "Quote Details/ Customer Details in Quote Cart section",
            "keywords": ["quote details", "customer details"],
            "description": "The solution must enable the user to view the relevant quote and customer details in the cart summary.",
            "criteria": "AC1: The system shall ensure that 'Quote Details' /'Customer details' section is collapsed by default.\nAC2: The system shall ensure that 'Quote Details' section is prepopulated with below details for New Customer.\nAC3: The system supports name: FirstName LastName.\nAC4: The system supports quote number: XXXXXXXXXXXXXX.\nAC5: The system shall ensure that 'Customer Details' section is prepopulated with below details for Existing Customer.\nAC6: The system supports active Services : XX.\nAC7: The system supports approved Services : XXX.\nAC8: Equipment Limit Remaining : $ X,XXX.\nAC9: Credit CHeck Status: Approved (as per status).", 
            "system": "Siebel",
        },
        {
            "ref": "FR024",
            "category": "Quote - Product and plan selection",
            "title": "Plan selection and Add to cart in Quote Add PLP page",
            "keywords": ["new services", "add to cart", "plan name"],
            "description": "The solution must enable the user to select a plan and add it to the cart from the quote product-list page.",
            "criteria": "AC1: The system shall ensure that 'New Services' section is collapsed by default.\nAC2: The system shall ensure that 'No new services added' text displayed if no plan is added to cart.\nAC3: The system shall ensure that The agent can add only one fixed plan to cart and add to cart by clicking 'Add to cart' button.\nAC4: The system shall ensure that Plan name with rate $XX.XX is displayed along with proposition name.",
            "system": "Siebel",
        },
        {
            "ref": "FR025",
            "category": "Quote - Product and plan selection",
            "title": "Modem selection and Add to cart in Quote Add PLP page",
            "keywords": ["modem selection", "add to cart", "modem name"],
            "description": "The solution must enable the user to select a modem and add it to the cart from the quote product-list page.",
            "criteria": "AC1: Modem Selection for FWA- Tech Feasibility.\nAC2: The system shall ensure that The agent can add only one Modem to Cart and add to cart by clicking 'Add to cart' button.\nAC3: The system shall ensure that Modem name with rate $XX.XX is displayed next to Plan in Cart.",
            "system": "Siebel",
        },
        {
            "ref": "FR026",
            "category": "Quote - Cart and quote creation",
            "title": "option to add NDC charge in Cart in Quote Add PLP page",
            "keywords": ["new development charge", "ndc charge", "add full month to next bill"],
            "description": "The solution must allow the user to add the new development charge to the cart when applicable and present the associated payment option.",
            "criteria": "AC1: The system shall ensure that popup should display with new development charges and payment option after click the 'Add to cart' button for order with development charges Payment option - Add full month to next bill and pay as term radio button and agent should allowed to select any option terms drop down is available with 12,24 and 36 months.\nAC2: The system shall ensure that 'New Development Charge' is displayed with charge $XX.XX next to Modem in Cart.\nAC3: The system shall ensure that popup for new development charges should not be displayed for Vision Technology while clicking 'Add to cart' button.\nAC4: The system shall ensure that popup for new development charges should not be displayed for FWA Technology while clicking 'Add to cart' button.",
            "system": "Siebel",
        },
    ]

    for template in requirement_templates:
        if any(keyword in low for keyword in template["keywords"]):
            append_unique(requirements, {
                "ref": template["ref"],
                "category": template["category"],
                "title": template["title"],
                "description": template["description"],
                "criteria": template["criteria"],
                "sme": "",
                "moscow": "Must",
                "system": template["system"],
                "ref_docs": "",
            })

    if not requirements:
        requirements.append({
            "ref": "FR001",
            "category": "Quote",
            "title": "Journey validation and system behaviour",
            "description": "The solution must apply the journey validation and business rules described in the attached TOM for the selected customer and journey stage.",
            "criteria": "AC1: The system evaluates the conditions defined in the TOM before proceeding to the next step.\nAC2: The system displays the associated status, message, or result for the action or validation outcome.",
            "sme": "",
            "moscow": "Must",
            "system": "Journey Platform",
            "ref_docs": "",
        })

    return requirements


def copy_reference_styles(template_ws, ws):
    for row in template_ws.iter_rows(min_row=1, max_row=min(2, template_ws.max_row), min_col=1, max_col=9):
        for cell in row:
            target = ws.cell(row=cell.row, column=cell.column)
            target.value = cell.value
            target.font = cell.font.copy()
            target.fill = cell.fill.copy()
            target.border = cell.border.copy()
            target.alignment = cell.alignment.copy()
            target.number_format = cell.number_format
    header = template_ws[4]
    for idx, cell in enumerate(header, start=1):
        target = ws.cell(row=4, column=idx)
        target.value = cell.value
        target.font = cell.font.copy()
        target.fill = cell.fill.copy()
        target.border = cell.border.copy()
        target.alignment = cell.alignment.copy()
        target.number_format = cell.number_format
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False


def style_headers(ws):
    for cell in ws[4]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="2F75B5")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="thin", color="D9E2F3"))


def apply_table_style(ws):
    for row in ws.iter_rows(min_row=5, min_col=1, max_col=9):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(
                left=Side(style="thin", color="D9E2F3"),
                right=Side(style="thin", color="D9E2F3"),
                top=Side(style="thin", color="D9E2F3"),
                bottom=Side(style="thin", color="D9E2F3"),
            )
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:I{ws.max_row}"
    for col_letter, width in {"A": 12, "B": 30, "C": 34, "D": 62, "E": 96, "F": 20, "G": 12, "H": 14, "I": 42}.items():
        ws.column_dimensions[col_letter].width = width


def build_workbook_from_template(template_wb, journey_name, requirements):
    wb = Workbook()
    ws = wb.active
    ws.title = "4. Functional Requirements"
    ws.sheet_view.showGridLines = False

    if "4. Functional Requirements" in template_wb.sheetnames:
        template_ws = template_wb["4. Functional Requirements"]
        title_text = template_ws["A1"].value or f"Business Requirements Specification - {journey_name}"
        subtitle_text = template_ws["A2"].value or "Functional requirements derived from TOM - Testcase Verification"
        ws.merge_cells("A1:I1")
        ws["A1"] = title_text
        ws["A1"].font = template_ws["A1"].font.copy()
        ws["A1"].alignment = template_ws["A1"].alignment.copy()
        ws["A1"].fill = template_ws["A1"].fill.copy()
        ws.merge_cells("A2:I2")
        ws["A2"] = subtitle_text
        ws["A2"].font = template_ws["A2"].font.copy()
        ws["A2"].alignment = template_ws["A2"].alignment.copy()
        ws["A2"].fill = template_ws["A2"].fill.copy()

        for idx, cell in enumerate(template_ws[4], start=1):
            target = ws.cell(row=4, column=idx)
            target.value = cell.value
            target.font = cell.font.copy()
            target.fill = cell.fill.copy()
            target.alignment = cell.alignment.copy()
            target.border = cell.border.copy()
    else:
        ws.merge_cells("A1:I1")
        ws["A1"] = f"Business Requirements Specification - {journey_name}"
        ws["A1"].font = Font(color="FFFFFF", bold=True, size=12)
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws["A1"].fill = PatternFill("solid", fgColor="1F3B6D")
        ws.merge_cells("A2:I2")
        ws["A2"] = "Functional requirements derived from TOM - Testcase Verification"
        ws["A2"].font = Font(color="000000", bold=True, size=10)
        ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws["A2"].fill = PatternFill("solid", fgColor="D9EAF7")
        for idx, header in enumerate(HEADER_ROW, start=1):
            ws.cell(row=4, column=idx, value=header)
        style_headers(ws)

    for req in requirements:
        ws.append([
            req["ref"],
            req["category"],
            req["title"],
            req["description"],
            req["criteria"],
            req["sme"],
            req["moscow"],
            req["system"],
            req["ref_docs"],
        ])

    apply_table_style(ws)

    if "1. Scope" in template_wb.sheetnames:
        wb.create_sheet("1. Scope")
        template_scope = template_wb["1. Scope"]
        scope_sheet = wb["1. Scope"]
        for row in template_scope.iter_rows():
            for cell in row:
                scope_sheet[cell.coordinate].value = cell.value
                if cell.has_style:
                    scope_sheet[cell.coordinate].font = cell.font.copy()
                    scope_sheet[cell.coordinate].fill = cell.fill.copy()
                    scope_sheet[cell.coordinate].border = cell.border.copy()
                    scope_sheet[cell.coordinate].alignment = cell.alignment.copy()
                    scope_sheet[cell.coordinate].number_format = cell.number_format
    else:
        scope_ws = wb.create_sheet("1. Scope")
        scope_ws["A1"] = "In Scope"
        scope_ws["B1"] = "Out of Scope"
        scope_ws["C1"] = "Assumptions"

    if "Image" in template_wb.sheetnames:
        template_image = template_wb["Image"]
        image_ws = wb.create_sheet("Image")
        for row in template_image.iter_rows():
            for cell in row:
                image_ws[cell.coordinate].value = cell.value
                if cell.has_style:
                    image_ws[cell.coordinate].font = cell.font.copy()
                    image_ws[cell.coordinate].fill = cell.fill.copy()
                    image_ws[cell.coordinate].border = cell.border.copy()
                    image_ws[cell.coordinate].alignment = cell.alignment.copy()
                    image_ws[cell.coordinate].number_format = cell.number_format
    else:
        wb.create_sheet("Image")

    if "Ref" in template_wb.sheetnames:
        template_ref = template_wb["Ref"]
        ref_ws = wb.create_sheet("Ref")
        for row in template_ref.iter_rows():
            for cell in row:
                ref_ws[cell.coordinate].value = cell.value
                if cell.has_style:
                    ref_ws[cell.coordinate].font = cell.font.copy()
                    ref_ws[cell.coordinate].fill = cell.fill.copy()
                    ref_ws[cell.coordinate].border = cell.border.copy()
                    ref_ws[cell.coordinate].alignment = cell.alignment.copy()
                    ref_ws[cell.coordinate].number_format = cell.number_format
    else:
        ref_ws = wb.create_sheet("Ref")
        ref_ws["A1"] = "Reference"

    return wb


def create_brs_workbook(journey_name: str, tom_path: Path, output_path: Path):
    template_path = find_reference_template(BASE_DIR)
    template_wb = load_workbook(template_path, data_only=False) if template_path else Workbook()
    tom_wb = load_workbook(tom_path, data_only=True)
    requirements = build_requirement_rows(tom_wb)
    wb = build_workbook_from_template(template_wb, journey_name, requirements)

    try:
        wb.save(output_path)
        return output_path
    except PermissionError:
        fallback = output_path.with_name(f"{output_path.stem}_{uuid.uuid4().hex[:8]}.xlsx")
        wb.save(fallback)
        return fallback


if __name__ == "__main__":
    tom_path = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_TOM
    if not tom_path.exists():
        raise FileNotFoundError(f"TOM workbook not found: {tom_path}")

    prompt_text = load_prompt(PROMPT_PATH)
    journey_name = infer_journey_name(tom_path)
    output_name = f"BRS_Generated_{journey_name}.xlsx"
    output_path = BASE_DIR / output_name
    result = create_brs_workbook(journey_name, tom_path, output_path)
    print(f"BRS output created: {result}")
    print(f"Prompt loaded: {len(prompt_text)} characters")
