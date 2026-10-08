from openpyxl import load_workbook

path = r"U:/Users/hm616288/Downloads/BRS-Agent/Updated BRS 5&6 23_09.xlsx"
wb = load_workbook(path, data_only=True)
print("SHEETS:", wb.sheetnames)
for s in wb.sheetnames:
    ws = wb[s]
    print(f"\n### {s} rows={ws.max_row} cols={ws.max_column}")
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12), values_only=True):
        print(row)
