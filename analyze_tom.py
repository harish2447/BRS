from openpyxl import load_workbook

path = "U:/Users/hm616288/Downloads/BRS-Agent/ImagineR3.7_SSJSFJ_J5and6_Oct Release_TOM_V0.9.xlsx"
wb = load_workbook(path, data_only=True)
print('SHEETS:', wb.sheetnames)
for s in wb.sheetnames:
    ws = wb[s]
    print(f'\n### {s} max_row={ws.max_row} max_col={ws.max_column}')
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 8), values_only=True):
        print(row)
